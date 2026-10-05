#!/usr/bin/env python3
"""
Score returned annotation workbooks.

Reads the completed Pass 1 and Pass 2 files and produces:

  1. Inter-annotator agreement (Fleiss' kappa, Krippendorff's alpha, pairwise
     Cohen's kappa) for the binary verdict and for each of the 17 categories.
  2. Agreement between the annotator majority and the dataset's labels. This is
     the number that answers the reviewer question "how do we know these labels
     are right?".
  3. A per-scenario error table: where the majority disagrees with the dataset.
  4. Data quality flag counts and reasoning ratings from Pass 2.

Calibration rows are excluded from all statistics by default.

Usage:
    python score_annotations.py --dir annotation_study
    python score_annotations.py --dir annotation_study --include-calibration
    python score_annotations.py --dir annotation_study --out results/
"""

import argparse
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

RISK_CATEGORIES = [
    "Allergy & Adverse Drug Reaction Risk", "Drug-Drug Interaction Risk",
    "Drug-Food Interaction Risk", "Dosage & Toxicity Risk",
    "Renal Impairment Risk", "Hepatic Impairment Risk",
    "Cardiac Impairment Risk", "Respiratory Impairment Risk",
    "Bleeding Risk", "Infection Risk", "Pregnancy & Breastfeeding Risk",
    "Alcohol Use Risk", "Tobacco Use Risk", "Substance Use Risk",
    "Caffeine Intake Risk", "Weight/BMI Risk", "Age Risk",
]

SHORT = {
    "Allergy & Adverse Drug Reaction Risk": "Allergy/ADR",
    "Drug-Drug Interaction Risk": "Drug-Drug",
    "Drug-Food Interaction Risk": "Drug-Food",
    "Dosage & Toxicity Risk": "Dose/Tox",
    "Renal Impairment Risk": "Renal",
    "Hepatic Impairment Risk": "Hepatic",
    "Cardiac Impairment Risk": "Cardiac",
    "Respiratory Impairment Risk": "Respiratory",
    "Bleeding Risk": "Bleeding",
    "Infection Risk": "Infection",
    "Pregnancy & Breastfeeding Risk": "Pregnancy",
    "Alcohol Use Risk": "Alcohol",
    "Tobacco Use Risk": "Tobacco",
    "Substance Use Risk": "Substance",
    "Caffeine Intake Risk": "Caffeine",
    "Weight/BMI Risk": "Weight/BMI",
    "Age Risk": "Age",
}

DATA_FLAGS = ["Implausible value", "Internal contradiction",
              "Missing needed field", "Unrealistic clinical setup",
              "Obsolete / unavailable drug"]


def as_bool(v):
    """Parse a CSV boolean. bool('False') is True, so do not use bool() on strings."""
    if isinstance(v, bool):
        return v
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return False
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return bool(int(v))
    s = str(v).strip().lower()
    if s in ("true", "t", "yes", "y", "1"):
        return True
    if s in ("false", "f", "no", "n", "0", ""):
        return False
    raise ValueError(f"cannot read boolean {v!r}")


def category_code(v):
    """1 TRUE, 0 FALSE, -1 missing or UNSURE."""
    if v is None:
        return -1
    s = str(v).strip().upper()
    if s in ("TRUE", "T", "YES", "Y", "1"):
        return 1
    if s in ("FALSE", "F", "NO", "N", "0"):
        return 0
    return -1


def confidence_filled(v):
    if v is None:
        return False
    s = str(v).strip()
    return s in {"1", "2", "3", "4", "5"} or s in {"1.0", "2.0", "3.0", "4.0", "5.0"}


# ==============================================================================
# Agreement statistics
# ==============================================================================

def fleiss_kappa(counts):
    """counts: (n_items, n_categories), each row summing to n_raters."""
    counts = np.asarray(counts, dtype=float)
    N, K = counts.shape
    n = counts.sum(axis=1)
    if not np.allclose(n, n[0]) or n[0] < 2:
        return float("nan")
    n = n[0]
    P_i = ((counts ** 2).sum(axis=1) - n) / (n * (n - 1))
    P_bar = P_i.mean()
    p_j = counts.sum(axis=0) / (N * n)
    P_e = (p_j ** 2).sum()
    denom = 1 - P_e
    return float((P_bar - P_e) / denom) if abs(denom) > 1e-12 else float("nan")


def krippendorff_alpha(labels):
    """Nominal alpha. labels: (n_items, n_raters), missing = -1."""
    labels = np.asarray(labels)
    values = sorted({int(v) for v in labels.flatten() if v >= 0})
    if len(values) < 2:
        return float("nan")
    coincidence = {(a, b): 0.0 for a in values for b in values}
    for row in labels:
        present = [int(v) for v in row if v >= 0]
        m = len(present)
        if m < 2:
            continue
        cnt = Counter(present)
        for a in values:
            for b in values:
                if a == b:
                    coincidence[(a, b)] += cnt[a] * (cnt[a] - 1) / (m - 1)
                else:
                    coincidence[(a, b)] += cnt[a] * cnt[b] / (m - 1)
    n_c = {v: sum(coincidence[(v, w)] for w in values) for v in values}
    n_total = sum(n_c.values())
    if n_total < 2:
        return float("nan")
    D_o = sum(coincidence[(a, b)] for a in values for b in values if a != b)
    D_e = sum(n_c[a] * n_c[b] for a in values for b in values if a != b) / (n_total - 1)
    return float(1 - D_o / D_e) if D_e > 0 else float("nan")


def cohen_kappa(a, b):
    a, b = np.asarray(a), np.asarray(b)
    mask = (a >= 0) & (b >= 0)
    a, b = a[mask], b[mask]
    if len(a) == 0:
        return float("nan")
    po = (a == b).mean()
    cats = sorted(set(a.tolist()) | set(b.tolist()))
    pe = sum((a == c).mean() * (b == c).mean() for c in cats)
    return float((po - pe) / (1 - pe)) if abs(1 - pe) > 1e-12 else float("nan")


def bootstrap_ci(fn, data, n_boot=2000, seed=0):
    rng = np.random.default_rng(seed)
    n = len(data)
    vals = []
    for _ in range(n_boot):
        idx = rng.integers(0, n, size=n)
        try:
            v = fn(data[idx])
            if not np.isnan(v):
                vals.append(v)
        except Exception:
            continue
    if not vals:
        return (float("nan"), float("nan"))
    return (float(np.percentile(vals, 2.5)), float(np.percentile(vals, 97.5)))


def band(k):
    if np.isnan(k):
        return "n/a"
    if k < 0:    return "poor"
    if k < 0.21: return "slight"
    if k < 0.41: return "fair"
    if k < 0.61: return "moderate"
    if k < 0.81: return "substantial"
    return "almost perfect"


# ==============================================================================
# Reading workbooks
# ==============================================================================

def _header_map(ws):
    found = {}
    for c in range(1, (ws.max_column or 0) + 1):
        v = ws.cell(row=1, column=c).value
        if v is not None and str(v).strip():
            found[str(v).strip()] = c
    return found


def _annotator_from_workbook(wb, path):
    if "Read First" in wb.sheetnames:
        ws = wb["Read First"]
        for r in range(1, 8):
            v = ws.cell(row=r, column=1).value
            if isinstance(v, str) and v.startswith("Annotator:"):
                return v.split(":", 1)[1].strip()
    return path.name.split("_")[0]


def read_pass1(path):
    """Returns {row_number: record}. A row is complete only if Confidence is set.

    Blank category cells are missing, not FALSE. Pre-filled FALSE is FALSE.
    """
    wb = load_workbook(path, data_only=True)
    if "Annotation" not in wb.sheetnames:
        raise ValueError(f"{path.name}: no 'Annotation' sheet")
    ws = wb["Annotation"]
    headers = _header_map(ws)
    missing = [h for h in ("Row", "Scenario ID", "Confidence 1-5") if h not in headers]
    missing += [SHORT[c] for c in RISK_CATEGORIES if SHORT[c] not in headers]
    if missing:
        raise ValueError(f"{path.name}: missing columns {missing}")

    overall_col = headers.get("Overall impression")
    comment_col = headers.get("Comments")
    out = {}
    for r in range(2, ws.max_row + 1):
        row_no = ws.cell(row=r, column=headers["Row"]).value
        if row_no is None:
            continue
        conf = ws.cell(row=r, column=headers["Confidence 1-5"]).value
        done = confidence_filled(conf)
        cats = {}
        for cat in RISK_CATEGORIES:
            raw = ws.cell(row=r, column=headers[SHORT[cat]]).value
            cats[cat] = category_code(raw) if done else -1
        impression = None
        if overall_col:
            raw_imp = ws.cell(row=r, column=overall_col).value
            if raw_imp is not None and str(raw_imp).strip():
                impression = str(raw_imp).strip().upper()
        out[int(row_no)] = {
            "categories": cats,
            "complete": done,
            "confidence": conf if done else None,
            "impression": impression if done else None,
            "comment": (ws.cell(row=r, column=comment_col).value if comment_col else None),
            "scenario_id": ws.cell(row=r, column=headers["Scenario ID"]).value,
        }
    return out, _annotator_from_workbook(wb, path)


def read_pass2(path):
    wb = load_workbook(path, data_only=True)
    if "Review" not in wb.sheetnames:
        return {}
    ws = wb["Review"]
    out = {}
    for r in range(2, ws.max_row + 1):
        row_no = ws.cell(row=r, column=1).value
        if row_no is None:
            continue
        flags = {}
        for k, name in enumerate(DATA_FLAGS):
            v = ws.cell(row=r, column=10 + k).value
            flags[name] = bool(v and str(v).strip().upper() == "X")
        out[int(row_no)] = {
            "agree": (str(ws.cell(row=r, column=7).value).strip().upper()
                      if ws.cell(row=r, column=7).value else None),
            "accuracy": ws.cell(row=r, column=8).value,
            "completeness": ws.cell(row=r, column=9).value,
            "flags": flags,
            "comment": ws.cell(row=r, column=10 + len(DATA_FLAGS)).value,
        }
    return out


# ==============================================================================
# Main analysis
# ==============================================================================

def main():
    ap = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dir", default="annotation_study")
    ap.add_argument("--out", default=None)
    ap.add_argument("--include-calibration", action="store_true")
    ap.add_argument("--bootstrap", type=int, default=2000)
    args = ap.parse_args()

    d = Path(args.dir)
    outdir = Path(args.out) if args.out else d / "results"
    outdir.mkdir(parents=True, exist_ok=True)

    key_path = d / "coordinator_only" / "_COORDINATOR_answer_key.csv"
    if not key_path.exists():
        key_path = d / "_COORDINATOR_answer_key.csv"
    if not key_path.exists():
        print(f"Answer key not found under {d}", file=sys.stderr)
        return 2
    key = pd.read_csv(key_path)

    manifest = {}
    mpath = d / "study_manifest.json"
    if mpath.exists():
        manifest = json.loads(mpath.read_text())
    n_calib = manifest.get("n_calibration", 0)

    p1_files = sorted(d.glob("*_pass1_blind.xlsx"))
    if not p1_files:
        print(f"No completed Pass 1 workbooks in {d}", file=sys.stderr)
        return 2

    annotations, annotators = {}, []
    for f in p1_files:
        try:
            rec, code = read_pass1(f)
            if code in annotations:
                print(f"  duplicate annotator code {code} from {f.name}", file=sys.stderr)
                return 2
            annotations[code] = rec
            annotators.append(code)
        except Exception as e:
            print(f"  could not read {f.name}: {e}", file=sys.stderr)

    print("=" * 74)
    print("ANNOTATION STUDY RESULTS")
    print("=" * 74)
    print(f"  annotators: {', '.join(annotators)} (n={len(annotators)})")

    rows = sorted(key["row"].tolist())
    if not args.include_calibration:
        rows = [r for r in rows if r > n_calib]
        print(f"  scenarios:  {len(rows)} (excluding {n_calib} calibration)")
    else:
        print(f"  scenarios:  {len(rows)} (including calibration)")

    # Completion: Confidence filled. Pre-filled FALSE with no Confidence is not an answer.
    print("\n  completion (Confidence filled):")
    for a in annotators:
        done = sum(1 for r in rows if annotations[a].get(r, {}).get("complete"))
        print(f"    {a}: {done}/{len(rows)}")

    keyrow = {int(r["row"]): r for _, r in key.iterrows()}

    # ---- Binary verdict ----
    print("\n" + "-" * 74)
    print("  BINARY VERDICT (safe = all 17 categories false)")
    print("-" * 74)

    headline = {}
    verdict = np.full((len(rows), len(annotators)), -1, dtype=int)
    for i, r in enumerate(rows):
        for j, a in enumerate(annotators):
            rec = annotations[a].get(r)
            if rec is None or not rec.get("complete"):
                verdict[i, j] = -1
                continue
            cats = rec["categories"]
            if any(v == 1 for v in cats.values()):
                verdict[i, j] = 0           # UNSAFE
            elif any(v == -1 for v in cats.values()):
                verdict[i, j] = -1          # UNSURE and nothing positive
            else:
                verdict[i, j] = 1           # SAFE

    valid = verdict.min(axis=1) >= 0
    vb = verdict[valid]
    if len(vb) >= 2:
        counts = np.column_stack([(vb == 0).sum(axis=1), (vb == 1).sum(axis=1)])
        k = fleiss_kappa(counts)
        klo, khi = bootstrap_ci(fleiss_kappa,
                                counts, args.bootstrap)
        alpha = krippendorff_alpha(verdict)
        alo, ahi = bootstrap_ci(krippendorff_alpha, verdict, args.bootstrap)
        print(f"    rows with all annotators responding: {len(vb)}/{len(rows)}")
        print(f"    Fleiss' kappa        {k:.3f}  [{klo:.3f}, {khi:.3f}]  {band(k)}")
        print(f"    Krippendorff's alpha {alpha:.3f}  [{alo:.3f}, {ahi:.3f}]")
        full = (vb == vb[:, [0]]).all(axis=1).mean()
        print(f"    unanimous            {full:.1%}")
        headline = {
            "fleiss_kappa": k,
            "fleiss_ci": [klo, khi],
            "krippendorff_alpha": alpha,
            "alpha_ci": [alo, ahi],
            "n_complete_rows": int(len(vb)),
            "unanimous": float(full),
            "band": band(k),
        }

        print("\n    pairwise Cohen's kappa:")
        print("      " + "".join(f"{a:>8}" for a in annotators))
        for i, a in enumerate(annotators):
            line = f"    {a:<4}"
            for j, b in enumerate(annotators):
                line += f"{1.0 if i == j else cohen_kappa(vb[:, i], vb[:, j]):>8.2f}"
            print(line)
    else:
        print("    not enough complete rows")

    # Stated impression, scored separately from the category-derived verdict.
    print("\n  stated overall impression (SAFE / UNSAFE):")
    impression = np.full_like(verdict, -1)
    mismatch = 0
    both = 0
    for i, r in enumerate(rows):
        for j, a in enumerate(annotators):
            rec = annotations[a].get(r) or {}
            if not rec.get("complete"):
                continue
            label = rec.get("impression")
            if label == "SAFE":
                impression[i, j] = 1
            elif label == "UNSAFE":
                impression[i, j] = 0
            if verdict[i, j] >= 0 and impression[i, j] >= 0:
                both += 1
                if verdict[i, j] != impression[i, j]:
                    mismatch += 1
    imp_valid = impression.min(axis=1) >= 0
    ib = impression[imp_valid]
    if len(ib) >= 2:
        icounts = np.column_stack([(ib == 0).sum(axis=1), (ib == 1).sum(axis=1)])
        ik = fleiss_kappa(icounts)
        print(f"    Fleiss' kappa        {ik:.3f}  {band(ik)}  (n={len(ib)})")
        headline["impression_fleiss_kappa"] = ik
    else:
        print("    not enough rows with an impression from every annotator")
    if both:
        print(f"    impression vs own categories: {mismatch}/{both} "
              f"({mismatch / both:.1%}) disagree")
        headline["impression_category_mismatches"] = mismatch
        headline["impression_category_compared"] = both

    # ---- Per category ----
    print("\n" + "-" * 74)
    print("  PER-CATEGORY AGREEMENT")
    print("-" * 74)
    print(f"    {'category':<14} {'n+':>4} {'kappa':>7} {'95% CI':>18}  {'alpha':>7}  band")

    per_cat = {}
    for cat in RISK_CATEGORIES:
        mat = np.full((len(rows), len(annotators)), -1, dtype=int)
        for i, r in enumerate(rows):
            for j, a in enumerate(annotators):
                rec = annotations[a].get(r)
                if rec and rec.get("complete"):
                    mat[i, j] = rec["categories"][cat]
        good = mat.min(axis=1) >= 0
        sub = mat[good]
        n_pos = int((sub == 1).any(axis=1).sum())
        # Alpha keeps rows with UNSURE or unfinished cells; kappa needs every rater.
        a_cat = krippendorff_alpha(mat)
        if len(sub) < 2 or n_pos == 0:
            print(f"    {SHORT[cat]:<14} {n_pos:>4} {'--':>7} {'':>18}  {a_cat:>7.3f}  no positives")
            per_cat[cat] = {"kappa": float("nan"), "alpha": a_cat, "n_positive": n_pos}
            continue
        counts = np.column_stack([(sub == 0).sum(axis=1), (sub == 1).sum(axis=1)])
        k = fleiss_kappa(counts)
        lo, hi = bootstrap_ci(fleiss_kappa, counts, min(args.bootstrap, 1000))
        alo_c, ahi_c = bootstrap_ci(krippendorff_alpha, mat, min(args.bootstrap, 1000))
        ci = f"[{lo:.2f}, {hi:.2f}]"
        print(f"    {SHORT[cat]:<14} {n_pos:>4} {k:>7.3f} {ci:>18}  {a_cat:>7.3f}  {band(k)}")
        per_cat[cat] = {"kappa": k, "ci": [lo, hi], "alpha": a_cat,
                        "alpha_ci": [alo_c, ahi_c], "n_positive": n_pos}

    ks = [v["kappa"] for v in per_cat.values() if not np.isnan(v["kappa"])]
    if ks:
        print(f"\n    macro kappa across {len(ks)} measurable categories: {np.mean(ks):.3f}")
    alphas = [v["alpha"] for v in per_cat.values() if not np.isnan(v.get("alpha", np.nan))]
    if alphas:
        print(f"    macro alpha across {len(alphas)} categories:          {np.mean(alphas):.3f}")

    # ---- Majority vs dataset ----
    print("\n" + "-" * 74)
    print("  ANNOTATOR MAJORITY vs DATASET LABELS")
    print("-" * 74)

    maj_v, ds_v, disagreements = [], [], []
    for i, r in enumerate(rows):
        if verdict[i].min() < 0:
            continue
        m = 1 if (verdict[i] == 1).sum() > len(annotators) / 2 else 0
        ds = 1 if as_bool(keyrow[r]["dataset_is_safe"]) else 0
        maj_v.append(m)
        ds_v.append(ds)
        if m != ds:
            disagreements.append({
                "row": r,
                "scenario_id": keyrow[r]["scenario_id"],
                "medication": keyrow[r].get("medication", ""),
                "dataset": "SAFE" if ds else "UNSAFE",
                "majority": "SAFE" if m else "UNSAFE",
                "vote_split": f"{(verdict[i]==1).sum()} safe / {(verdict[i]==0).sum()} unsafe",
            })

    if maj_v:
        maj_v, ds_v = np.array(maj_v), np.array(ds_v)
        k = cohen_kappa(maj_v, ds_v)
        agree = (maj_v == ds_v).mean()
        print(f"    scenarios compared:  {len(maj_v)}")
        print(f"    raw agreement:       {agree:.1%}")
        print(f"    Cohen's kappa:       {k:.3f}  {band(k)}")
        print(f"    disagreements:       {len(disagreements)}")
        print()
        print("    This is the number that answers 'how do we know the labels")
        print("    are right?'. Report it in the paper alongside the")
        print("    inter-annotator kappa above; the labels are defensible when")
        print("    this sits inside the annotators' own agreement envelope.")

        if disagreements:
            print(f"\n    scenarios where the majority disagrees with the dataset:")
            print(f"      {'row':>4} {'id':>8} {'dataset':>8} {'majority':>9}  vote")
            for dd in disagreements[:25]:
                print(f"      {dd['row']:>4} {str(dd['scenario_id']):>8} "
                      f"{dd['dataset']:>8} {dd['majority']:>9}  {dd['vote_split']}")
            if len(disagreements) > 25:
                print(f"      ... and {len(disagreements) - 25} more")

    # ---- Per-category majority vs dataset ----
    print("\n    per-category majority vs dataset:")
    print(f"      {'category':<14} {'kappa':>7}  {'FP':>4} {'FN':>4}")
    cat_vs_ds = {}
    for cat in RISK_CATEGORIES:
        m_list, d_list = [], []
        for i, r in enumerate(rows):
            votes = []
            for a in annotators:
                rec = annotations[a].get(r)
                if rec and rec["categories"][cat] >= 0:
                    votes.append(rec["categories"][cat])
            if len(votes) < len(annotators):
                continue
            m_list.append(1 if sum(votes) > len(votes) / 2 else 0)
            d_list.append(1 if as_bool(keyrow[r][cat]) else 0)
        if not m_list or (sum(m_list) == 0 and sum(d_list) == 0):
            continue
        m_arr, d_arr = np.array(m_list), np.array(d_list)
        k = cohen_kappa(m_arr, d_arr)
        fp = int(((d_arr == 1) & (m_arr == 0)).sum())   # dataset says yes, humans no
        fn = int(((d_arr == 0) & (m_arr == 1)).sum())   # humans say yes, dataset no
        cat_vs_ds[cat] = {"kappa": k, "dataset_only": fp, "humans_only": fn}
        print(f"      {SHORT[cat]:<14} {k:>7.3f}  {fp:>4} {fn:>4}")
    print("      FP = dataset marked TRUE, annotators did not")
    print("      FN = annotators marked TRUE, dataset did not")

    # ---- Pass 2 ----
    p2_files = sorted(d.glob("*_pass2_review.xlsx")) + sorted((d / "pass2_hold").glob("*_pass2_review.xlsx"))
    reviews = {}
    for f in p2_files:
        code = f.name.split("_")[0]
        try:
            r = read_pass2(f)
            if any(v["accuracy"] is not None or any(v["flags"].values())
                   for v in r.values()):
                reviews[code] = r
        except Exception:
            continue

    if reviews:
        print("\n" + "-" * 74)
        print("  PASS 2: REASONING AND DATA QUALITY")
        print("-" * 74)

        acc, comp = [], []
        flag_counts = Counter()
        flagged_rows = defaultdict(set)
        for code, rec in reviews.items():
            for r, v in rec.items():
                if not args.include_calibration and r <= n_calib:
                    continue
                if isinstance(v["accuracy"], (int, float)):
                    acc.append(v["accuracy"])
                if isinstance(v["completeness"], (int, float)):
                    comp.append(v["completeness"])
                for name, on in v["flags"].items():
                    if on:
                        flag_counts[name] += 1
                        flagged_rows[name].add(r)

        agree_counts = Counter()
        for code, rec in reviews.items():
            for r, v in rec.items():
                if not args.include_calibration and r <= n_calib:
                    continue
                if v.get("agree"):
                    agree_counts[v["agree"]] += 1
        if agree_counts:
            total_ag = sum(agree_counts.values())
            print("    agree with dataset verdict (Pass 2, after seeing reasoning):")
            for name in ("AGREE", "DISAGREE", "UNSURE"):
                n = agree_counts[name]
                if n:
                    print(f"      {name:<10} {n:>4}  ({n / total_ag:.1%})")

        if acc:
            print(f"    reasoning accuracy      mean {np.mean(acc):.2f}  "
                  f"median {np.median(acc):.1f}  n={len(acc)}")
            low = sum(1 for x in acc if x <= 2)
            print(f"      rated 2 or below:     {low} ({100*low/len(acc):.1f}%)")
        if comp:
            print(f"    reasoning completeness  mean {np.mean(comp):.2f}  "
                  f"median {np.median(comp):.1f}  n={len(comp)}")

        if flag_counts:
            print("\n    data quality flags (total marks / distinct scenarios):")
            for name in DATA_FLAGS:
                n = flag_counts[name]
                if n:
                    print(f"      {name:<28} {n:>4} / {len(flagged_rows[name]):>3}")
            rowsf = []
            for name, rs in flagged_rows.items():
                for r in sorted(rs):
                    rowsf.append({"row": r,
                                  "scenario_id": keyrow.get(r, {}).get("scenario_id", ""),
                                  "flag": name})
            pd.DataFrame(rowsf).to_csv(outdir / "data_quality_flags.csv", index=False)
    else:
        print("\n  (no completed Pass 2 workbooks found yet)")

    # ---- Save ----
    if disagreements:
        adj = pd.DataFrame(disagreements)
        adj["adjudication"] = ""
        adj["adjudicator_note"] = ""
        adj.to_csv(outdir / "verdict_disagreements.csv", index=False)

    results = {
        "annotators": annotators,
        "n_scenarios": len(rows),
        "excluded_calibration": (not args.include_calibration),
        "binary_verdict": headline,
        "per_category_agreement": {k: {kk: (None if isinstance(vv, float) and np.isnan(vv) else vv)
                                       for kk, vv in v.items()}
                                   for k, v in per_cat.items()},
        "per_category_vs_dataset": cat_vs_ds,
        "n_verdict_disagreements": len(disagreements),
    }
    with open(outdir / "agreement_results.json", "w") as f:
        json.dump(results, f, indent=2, default=str)

    print(f"\n  wrote results to {outdir}/")
    print()
    return 0


if __name__ == "__main__":
    sys.exit(main())
